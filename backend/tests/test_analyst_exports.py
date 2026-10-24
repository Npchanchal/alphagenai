"""CiteAlpha Analyst exports: per-cell citations, source footers, receipt = run log (US-A07, A4 gate)."""

import base64
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.data.seed import reset_data
from app.main import app
from app.services.agent import exports
from app.services.citations import lookup_citation
from tests.test_analyst_projects import MEMO, MODEL, _project, _upload

client = TestClient(app)
DEMO = {"X-API-Key": "intellens-demo"}
OTHER = {"X-API-Key": "intellens-onestop"}
RETAIL = {"X-API-Key": "intellens-retail"}


@pytest.fixture(autouse=True)
def _tmp_store(tmp_path, monkeypatch):
    monkeypatch.setenv("INTELLENS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("INTELLENS_PUBLIC_URL", "https://citealpha.example")
    reset_data()


def _run(skill="peer_delivery_grid", company="infy", **extra):
    r = client.post("/api/analyst/run", json={"skill_id": skill, "company_id": company, **extra}, headers=DEMO)
    assert r.status_code == 200, r.text
    return r.json()


def _export(run_id, fmt, headers=DEMO):
    return client.get(f"/api/analyst/runs/{run_id}/export?format={fmt}", headers=headers)


def test_peer_grid_xlsx_every_sourced_cell_has_quote_and_citation_link():
    run = _run()
    r = _export(run["run_id"], "xlsx")
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml")
    assert 'filename="citealpha-peer-delivery-grid-infy-' in r.headers["content-disposition"]
    wb = load_workbook(io.BytesIO(r.content))
    by_n = {c["n"]: c for c in run["citations"]}

    brief = wb["Brief"]
    header_row = next(c.row for c in brief["B"] if c.value == "Statement")
    statements = [row for row in brief.iter_rows(min_row=header_row + 1) if row[2].value]
    assert statements
    for section, statement, refs in statements:
        assert statement.comment is not None, statement.value
        first = int(refs.value.strip("[").split("]")[0])
        assert by_n[first]["quote"].split()[0] in statement.comment.text

    tables = [ws for ws in wb if ws.title not in ("Brief", "Sources", "Evidence receipt")]
    assert tables, "Peer Delivery Grid exports a table sheet"
    cells = [c for ws in tables for row in ws.iter_rows(min_row=2) for c in row[:-1]]
    assert cells
    for c in cells:
        assert c.comment is not None and "“" in c.comment.text
        assert c.hyperlink is not None and "/c/cite_" in c.hyperlink.target
        assert lookup_citation(c.hyperlink.target.rsplit("/c/", 1)[1]) is not None


def test_xlsx_receipt_matches_saved_run_log():
    run = _run("results_prep")
    saved = client.get(f"/api/analyst/runs/{run['run_id']}", headers=DEMO).json()
    wb = load_workbook(io.BytesIO(_export(run["run_id"], "xlsx").content))
    ws = wb["Evidence receipt"]
    fields = {r[0].value: r[1].value for r in ws.iter_rows(max_row=10)}
    assert fields["Run"] == saved["evidence_receipt"]["run_id"]
    assert fields["Skill"] == f"{saved['skill']['id']} v{saved['skill']['version']}"
    head = next(c.row for c in ws["G"] if c.value == "Quote SHA-256")
    hashes = [ws.cell(row=r, column=7).value for r in range(head + 1, head + 1 + len(saved["evidence_receipt"]["sources"]))]
    assert hashes == [s["quote_sha256"] for s in saved["evidence_receipt"]["sources"]]
    assert ws.max_row > head + len(hashes)  # tool calls follow

    src = wb["Sources"]
    assert src.max_row - 1 == len(saved["citations"])


def test_pptx_has_source_footers_and_receipt():
    run = _run("results_prep")
    r = _export(run["run_id"], "pptx")
    assert r.status_code == 200
    prs = Presentation(io.BytesIO(r.content))
    texts = ["\n".join(s.text_frame.text for s in slide.shapes) for slide in prs.slides]
    section_slides = [t for t in texts if "• " in t and "What we could not find" not in t]
    assert section_slides and all("Sources: [" in t for t in section_slides)
    links = [r_.hyperlink.address for slide in prs.slides for s in slide.shapes
             for p in s.text_frame.paragraphs for r_ in p.runs if r_.hyperlink.address]
    assert any("/c/cite_" in a for a in links)
    assert any(t.startswith("Evidence receipt") for t in texts)
    joined = "\n".join(texts)
    assert all(s["quote_sha256"] in joined for s in run["evidence_receipt"]["sources"])
    assert "Not investment advice" in texts[0]


def test_pdf_memo_and_markdown_carry_receipt():
    run = _run("company_ramp")
    pdf = _export(run["run_id"], "pdf")
    assert pdf.status_code == 200 and pdf.content[:4] == b"%PDF"
    text = "\n".join(p.extract_text() for p in PdfReader(io.BytesIO(pdf.content)).pages)
    assert "EVIDENCE RECEIPT" in text and run["run_id"] in text and "SOURCES" in text
    md = _export(run["run_id"], "md")
    assert md.headers["content-type"].startswith("text/markdown")
    assert run["evidence_receipt"]["sources"][0]["quote_sha256"] in md.text


def test_export_access_and_format():
    run = _run()
    assert _export(run["run_id"], "docx").status_code == 422
    assert _export(run["run_id"], "xlsx", headers=OTHER).status_code == 404
    assert _export(run["run_id"], "xlsx", headers=RETAIL).status_code == 403
    assert _export("run_missing", "xlsx").status_code == 404


def test_upload_cites_have_no_public_link():
    pid = _project()
    assert _upload(pid, "model.xlsx", MODEL).status_code == 200
    assert _upload(pid, "memo.docx", MEMO).status_code == 200
    run = _run("diligence_room", project_id=pid)
    wb = load_workbook(io.BytesIO(_export(run["run_id"], "xlsx").content))
    upload_ns = {c["n"] for c in run["citations"] if c.get("kind") == "upload"}
    assert upload_ns
    rows = [row for row in wb["Sources"].iter_rows(min_row=2) if row[0].value in upload_ns]
    assert len(rows) == len(upload_ns)
    for row in rows:
        assert row[6].value == "your upload" and row[6].hyperlink is None
        assert row[7].value is None
    pdf = _export(run["run_id"], "pdf")
    assert b"upload://" not in pdf.content


def _stub_run(text):
    return {
        "run_id": "run_stub", "skill": {"id": "stub", "version": "0", "title": "Stub"}, "company_id": "x",
        "company_name": "Stub Co", "as_of": "2026-01-01", "generated_at": "2026-01-01T00:00:00Z",
        "headline": None, "not_found": [], "disclaimer": "Stub.", "verifier": {"removed": 0},
        "sections": [{"id": "s", "title": "S", "claims": [{"text": text, "cites": [1], "kind": "quoted"}]}],
        "citations": [{"n": 1, "kind": "upload", "title": "f.csv (line 2) — your upload", "quote": text,
                       "source_url": "upload://u#line 2"}],
        "evidence_receipt": {"run_id": "run_stub", "sources": [], "tool_calls": []},
    }


def test_xlsx_never_writes_formulas_from_quotes():
    wb = load_workbook(io.BytesIO(exports.to_xlsx(_stub_run('=HYPERLINK("http://x","y")'), "https://b")))
    cells = [c for ws in wb for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    assert cells and all(c.data_type == "s" for c in cells)


def test_pdf_handles_non_ascii():
    out = exports.to_pdf(_stub_run("Capex of ₹1,200 cr — “on track” for FY27…"), "https://b")
    assert "Rs 1,200 cr - \"on track\"" in PdfReader(io.BytesIO(out)).pages[0].extract_text()
