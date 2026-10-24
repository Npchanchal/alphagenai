"""A desk's house style on Analyst exports: colours, font, firm / footer / classification, logo, PowerPoint template."""

import base64
import io
import struct
import zlib

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Inches

from app.data.seed import reset_data
from app.main import app

client = TestClient(app)
ADMIN = {"X-API-Key": "intellens-admin"}
DEMO = {"X-API-Key": "intellens-demo"}
OTHER = {"X-API-Key": "intellens-onestop"}
STYLE = {"firm_name": "Acme Securities", "primary_color": "#123456", "accent_color": "#AA0000", "font": "Georgia",
         "footer": "Acme Research · Mumbai", "classification": "Internal — not for distribution"}


@pytest.fixture(autouse=True)
def _tmp_store(tmp_path, monkeypatch):
    monkeypatch.setenv("INTELLENS_DATA_DIR", str(tmp_path))
    monkeypatch.setenv("INTELLENS_PUBLIC_URL", "https://citealpha.example")
    reset_data()


def _png() -> bytes:
    def chunk(t, d):
        return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xFFFFFFFF)

    raw = b"\x00\xff\x00\x00"
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0))
            + chunk(b"IDAT", zlib.compress(raw)) + chunk(b"IEND", b""))


def _template() -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(10), Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "HOUSE TEMPLATE COVER"
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _run():
    r = client.post("/api/analyst/run", json={"skill_id": "peer_delivery_grid", "company_id": "infy"}, headers=DEMO)
    assert r.status_code == 200, r.text
    return r.json()


def _texts(prs):
    return [p.text for s in prs.slides for sh in s.shapes if sh.has_text_frame for p in sh.text_frame.paragraphs]


def test_only_admin_edits_and_any_seat_reads():
    assert client.put("/api/analyst/house-style", json=STYLE, headers=DEMO).status_code == 403
    assert client.put("/api/analyst/house-style", json=STYLE, headers=ADMIN).status_code == 200
    got = client.get("/api/analyst/house-style", headers=DEMO).json()
    assert got["firm_name"] == "Acme Securities" and got["configured"] is True
    assert client.get("/api/analyst/house-style", headers=OTHER).json()["configured"] is False


def test_validation():
    assert client.put("/api/analyst/house-style", json={**STYLE, "primary_color": "red"}, headers=ADMIN).status_code == 422
    assert client.put("/api/analyst/house-style", json={**STYLE, "font": "Comic Sans"}, headers=ADMIN).status_code == 422
    gif = base64.b64encode(b"GIF89a" + b"\0" * 20).decode()
    assert client.put("/api/analyst/house-style", json={**STYLE, "logo_base64": gif}, headers=ADMIN).status_code == 422
    bad = base64.b64encode(b"not a zip").decode()
    r = client.post("/api/analyst/house-style/template", json={"filename": "x.pptx", "content_base64": bad}, headers=ADMIN)
    assert r.status_code == 422
    r = client.post("/api/analyst/house-style/template", json={"filename": "x.docx", "content_base64": bad}, headers=ADMIN)
    assert r.status_code == 422


def test_pptx_in_house_style_keeps_citations_and_receipt():
    client.put("/api/analyst/house-style", json={**STYLE, "logo_base64": base64.b64encode(_png()).decode()}, headers=ADMIN)
    run = _run()
    r = client.get(f"/api/analyst/runs/{run['run_id']}/export?format=pptx", headers=DEMO)
    prs = Presentation(io.BytesIO(r.content))
    texts = _texts(prs)
    assert any("Prepared for Acme Securities" in t for t in texts)
    assert any("Internal — not for distribution" in t for t in texts)
    assert any(t.startswith("Sources: ") for t in texts), "source footers are kept"
    assert any(t == "Evidence receipt" for t in texts)
    fonts = {r_.font.name for s in prs.slides for sh in s.shapes if sh.has_text_frame
             for p in sh.text_frame.paragraphs for r_ in p.runs}
    assert fonts == {"Georgia"}
    title = prs.slides[0].shapes[0].text_frame.paragraphs[0].runs[0]
    assert str(title.font.color.rgb) == "123456"
    assert any(sh.shape_type == 13 for sh in prs.slides[0].shapes), "logo picture on the slide"


def test_pptx_on_desk_template_uses_its_size_and_drops_sample_slides():
    client.put("/api/analyst/house-style", json=STYLE, headers=ADMIN)
    tpl = base64.b64encode(_template()).decode()
    r = client.post("/api/analyst/house-style/template", json={"filename": "acme.pptx", "content_base64": tpl}, headers=ADMIN)
    assert r.status_code == 200 and r.json()["has_template"] is True
    run = _run()
    prs = Presentation(io.BytesIO(client.get(f"/api/analyst/runs/{run['run_id']}/export?format=pptx", headers=DEMO).content))
    assert prs.slide_width == Inches(10)
    texts = _texts(prs)
    assert "HOUSE TEMPLATE COVER" not in texts
    assert any(t.startswith("Peer Delivery Grid") for t in texts)
    assert any(t == "Evidence receipt" for t in texts)
    assert client.delete("/api/analyst/house-style/template", headers=ADMIN).json()["has_template"] is False


def test_template_with_doctype_rejected():
    import zipfile

    buf = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(_template())) as src, zipfile.ZipFile(buf, "w") as dst:
        for info in src.infolist():
            data = src.read(info)
            if info.filename == "ppt/presentation.xml":
                data = data.replace(b"?>", b'?><!DOCTYPE x [<!ENTITY e SYSTEM "file:///etc/passwd">]>', 1)
            dst.writestr(info, data)
    r = client.post("/api/analyst/house-style/template",
                    json={"filename": "x.pptx", "content_base64": base64.b64encode(buf.getvalue()).decode()}, headers=ADMIN)
    assert r.status_code == 422


def test_xlsx_pdf_md_carry_house_lines():
    client.put("/api/analyst/house-style", json=STYLE, headers=ADMIN)
    run = _run()
    wb = load_workbook(io.BytesIO(client.get(f"/api/analyst/runs/{run['run_id']}/export?format=xlsx", headers=DEMO).content))
    brief = wb["Brief"]
    assert brief["A2"].value == "Prepared for Acme Securities"
    assert "Internal" in brief.oddFooter.center.text
    hdr = next(c for c in brief["B"] if c.value == "Statement")
    assert hdr.fill.fgColor.rgb.endswith("123456")
    md = client.get(f"/api/analyst/runs/{run['run_id']}/export?format=md", headers=DEMO).text
    assert md.startswith("> Prepared for Acme Securities") and "## Evidence receipt" in md
    pdf = client.get(f"/api/analyst/runs/{run['run_id']}/export?format=pdf", headers=DEMO)
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")


def test_style_is_per_org_and_reset_restores_default():
    client.put("/api/analyst/house-style", json=STYLE, headers=ADMIN)
    other_run = client.post("/api/analyst/run", json={"skill_id": "peer_delivery_grid", "company_id": "infy"}, headers=OTHER).json()
    md = client.get(f"/api/analyst/runs/{other_run['run_id']}/export?format=md", headers=OTHER).text
    assert "Acme" not in md
    client.delete("/api/analyst/house-style", headers=ADMIN)
    run = _run()
    md = client.get(f"/api/analyst/runs/{run['run_id']}/export?format=md", headers=DEMO).text
    assert "Acme" not in md


def test_preview_deck():
    client.put("/api/analyst/house-style", json=STYLE, headers=ADMIN)
    r = client.get("/api/analyst/house-style/preview", headers=DEMO)
    assert r.status_code == 200
    assert any("Acme Securities" in t for t in _texts(Presentation(io.BytesIO(r.content))))
