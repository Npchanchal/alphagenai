"""Compliance sign-off gate for CiteAlpha Analyst Skills and export layouts.

docs/compliance/ANALYST_SKILLS_SIGNOFF.md records the reviewed Skill versions and
export layout version. A changed Skill or layout fails here until it is re-reviewed
and the record updated. Every export, in the default layout and in a house style,
must keep the disclaimer, the data date, source quotes of 400 characters or less
and the Evidence Receipt; fixed layout copy carries no advice wording.
"""

import io
import re
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pptx import Presentation
from pypdf import PdfReader

from app.services.agent import exports, runtime, skills
from tests.test_analyst_agent import SKILL_COPY_BANNED

RECORD = Path(__file__).resolve().parents[2] / "docs" / "compliance" / "ANALYST_SKILLS_SIGNOFF.md"
BASE = "https://citealpha.example"
HOUSE = {"firm_name": "Acme Securities", "primary_color": "#123456", "accent_color": "#AA0000", "font": "Georgia",
         "footer": "Acme Research", "classification": "Internal", "logo": None, "template": None}
DISCLAIMER = "Not investment advice"


def _text(fmt: str, blob: bytes) -> str:
    if fmt == "xlsx":
        wb = load_workbook(io.BytesIO(blob))
        return "\n".join(wb.sheetnames + [str(c.value) for ws in wb for row in ws.iter_rows() for c in row
                                          if c.value is not None])
    if fmt == "pptx":
        prs = Presentation(io.BytesIO(blob))
        return "\n".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    if fmt == "pdf":
        return "\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(blob)).pages)
    return blob.decode("utf-8")


def test_signoff_record_covers_every_skill_version_and_layout():
    record = RECORD.read_text(encoding="utf-8")
    missing = [f"{sid} v{spec['version']}" for sid, spec in skills.load_skills().items()
               if not re.search(rf"\|\s*`{sid}`\s*\|\s*{re.escape(spec['version'])}\s*\|", record)]
    assert not missing, f"Re-review and add to {RECORD.name}: {missing}"
    assert f"Export layout version: **{exports.LAYOUT_VERSION}**" in record
    assert "Human approver" in record


@pytest.mark.parametrize("skill_id", sorted(skills.load_skills()))
def test_skill_spec_is_factual_and_has_a_disproof_step(skill_id):
    spec = skills.get_skill(skill_id)
    assert spec.get("disproof"), "every Skill states what would disprove its output"
    assert spec.get("description") and spec.get("audience")
    for sec in spec["sections"]:
        assert len(sec.get("title", "")) <= 80


@pytest.mark.parametrize("skill_id", sorted(skills.load_skills()))
def test_every_export_keeps_disclaimer_date_receipt_and_quote_limit(skill_id):
    run = runtime.run_skill(skill_id, "infy", use_llm=False)
    assert DISCLAIMER in run["disclaimer"]
    for style in (None, HOUSE):
        for fmt in exports.FORMATS:
            text = " ".join(_text(fmt, exports.render(run, fmt, BASE, style=style)).split())
            assert DISCLAIMER in text, (fmt, style is not None)
            assert f"Data as of {run['as_of']}" in text, fmt
            assert re.search(r"evidence receipt", text, re.I), fmt
            assert run["run_id"] in text, fmt
    for c in run["citations"]:
        assert len(exports._quote(c)) <= exports.MAX_QUOTE


def test_fixed_layout_copy_has_no_advice_wording():
    sample = {
        "run_id": "layout", "skill": {"id": "layout_check", "title": "Layout", "version": "1"},
        "company_id": "", "company_name": "Layout", "as_of": "2026-01-01", "generated_at": "",
        "disclaimer": "Factual summary of company disclosures with sources. Not investment advice.",
        "headline": "", "sections": [], "citations": [], "not_found": ["Nothing found"],
        "evidence_receipt": {"sources": [], "tool_calls": []},
    }
    for style in (None, HOUSE):
        for fmt in exports.FORMATS:
            text = _text(fmt, exports.render(sample, fmt, BASE, style=style))
            hits = [m.group(0) for pat in SKILL_COPY_BANNED for m in re.finditer(pat, text, re.I)]
            assert not hits, (fmt, hits)


def test_house_style_cannot_drop_citations():
    run = runtime.run_skill("peer_delivery_grid", "infy", use_llm=False)
    plain = _text("pptx", exports.render(run, "pptx", BASE))
    styled = _text("pptx", exports.render(run, "pptx", BASE, style=HOUSE))
    for c in run["citations"]:
        tag = f"[{c['n']}]"
        assert (tag in plain) == (tag in styled)
